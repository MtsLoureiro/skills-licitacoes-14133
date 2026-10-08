#!/usr/bin/env python3
"""Gera os documentos da pesquisa de preços a partir de <prefixo>.calculo.json.

Saídas (com --out PREFIXO):
  PREFIXO.mapa.csv           mapa comparativo (um preço por linha, com fonte PNCP)
  PREFIXO.mapa.xlsx          o mesmo em planilha (só se openpyxl estiver instalado)
  PREFIXO.memoria.md         memória de cálculo (método, critério de corte, passo a passo)
  PREFIXO.justificativa.md   minuta da "Justificativa da pesquisa de preços" (Markdown)

Os textos são MINUTAS: preencha os campos [ENTRE COLCHETES], revise e adapte à norma do seu
ente antes de juntar ao processo. O script não inventa fonte: só usa o que veio no cálculo.
"""
import argparse
import csv
import json
import sys
from datetime import date

COLUNAS = [
    ("n", "Nº"), ("data_resultado", "Data do resultado"), ("_valor", "Preço unitário (R$)"),
    ("unidade_fornecimento", "Unidade"), ("quantidade", "Quantidade"), ("fornecedor", "Fornecedor"),
    ("orgao", "Órgão"), ("uf", "UF"), ("modalidade", "Modalidade"), ("numero_controle_pncp", "Nº controle PNCP"),
    ("link_pncp", "Link PNCP"), ("alertas", "Alertas"),
]


def brl(v, casas=2):
    if v is None:
        return "-"
    s = f"{v:,.{casas}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def linhas_mapa(res):
    out = []
    ordenados = sorted(res["mantidos"], key=lambda r: r["_valor"])
    for i, r in enumerate(ordenados, 1):
        out.append({
            "n": i, "data_resultado": r.get("data_resultado") or "", "_valor": round(r["_valor"], 4),
            "unidade_fornecimento": r.get("unidade_fornecimento") or r.get("unidade_medida") or "",
            "quantidade": r.get("quantidade") if r.get("quantidade") is not None else "",
            "fornecedor": r.get("fornecedor") or "", "orgao": r.get("orgao") or "", "uf": r.get("uf") or "",
            "modalidade": r.get("modalidade") or "", "numero_controle_pncp": r.get("numero_controle_pncp") or "",
            "link_pncp": r.get("link_pncp") or "", "alertas": " | ".join(r.get("alertas", [])),
        })
    return out


def grava_csv(linhas, caminho):
    with open(caminho, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([t for _, t in COLUNAS])
        for l in linhas:
            w.writerow([l[k] for k, _ in COLUNAS])


def grava_xlsx(linhas, res, caminho):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("[info] openpyxl não instalado: pulei o .xlsx (pip install openpyxl)", file=sys.stderr)
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Mapa de Preços"
    ws.append([t for _, t in COLUNAS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for l in linhas:
        ws.append([l[k] for k, _ in COLUNAS])
    ult = len(linhas) + 1
    col = "C"
    ws.append([])
    ws.append(["", "Média", f"=AVERAGE({col}2:{col}{ult})"])
    ws.append(["", "Mediana", f"=MEDIAN({col}2:{col}{ult})"])
    ws.append(["", "Menor", f"=MIN({col}2:{col}{ult})"])
    ws.append(["", "Desvio-padrão", f"=STDEV({col}2:{col}{ult})"])
    ws.append(["", f"Preço estimado ({res['parametros']['metodo']})", res["preco_estimado"]])
    for i, larg in enumerate([5, 14, 16, 12, 12, 38, 38, 5, 22, 30, 52, 60]):
        ws.column_dimensions[chr(65 + i)].width = larg
    ws2 = wb.create_sheet("Excluídos")
    ws2.append(["Linha", "Motivo", "Preço", "Fornecedor", "Órgão", "Nº controle PNCP"])
    for r in res["excluidos"]:
        ws2.append([r.get("_linha"), r.get("_excluido"), r.get(res["parametros"]["campo_preco"]),
                    r.get("fornecedor"), r.get("orgao"), r.get("numero_controle_pncp")])
    wb.save(caminho)
    return True


def memoria_md(res, titulo):
    p, f, a0 = res["parametros"], res["estatisticas_finais"], res["estatisticas_antes_do_corte"]
    L = [f"# Memória de cálculo — {titulo}", "", f"Gerada em {date.today().strftime('%d/%m/%Y')} por `calcular_precos.py`.", "",
         "## 1. Parâmetros aplicados", "",
         f"- Método do preço estimado: **{p['metodo']}**",
         f"- Critério de corte de valores inconsistentes/inexequíveis/excessivamente elevados: **{p['outliers']}** — {res['criterio_corte']}",
         f"- Fator de atualização aplicado: {p['fator_atualizacao']:g}",
         f"- Mínimo de preços exigido: {p['min_precos']}",
         f"- Filtros: quantidade mín. {p['qtd_min']}, máx. {p['qtd_max']}, UF {p['uf'] or 'todas'}, palavras {p['palavras'] or '—'}", "",
         "## 2. Funil dos dados", "",
         f"- Registros de entrada: {res['entrada_total']}",
         f"- Após filtros objetivos e remoção de duplicatas: {res['apos_filtros_e_duplicatas']}",
         f"- Preços válidos após o corte: {f.get('n', 0)}", ""]
    if a0.get("n"):
        L += ["## 3. Estatísticas", "", "| | n | Média | Mediana | Menor | Maior | Desvio-padrão | CV |", "|---|---|---|---|---|---|---|---|"]
        for nome, s in (("Antes do corte", a0), ("Depois do corte", f)):
            if s.get("n"):
                cv = f"{s['coef_variacao']:.1%}" if s.get("coef_variacao") is not None else "-"
                L.append(f"| {nome} | {s['n']} | {brl(s['media'], 4)} | {brl(s['mediana'], 4)} | {brl(s['menor'], 4)} | {brl(s['maior'], 4)} | {brl(s['desvio_padrao'], 4)} | {cv} |")
        L.append("")
    L += ["## 4. Resultado", ""]
    if res["preco_estimado"] is not None:
        L.append(f"**Preço unitário estimado ({p['metodo']}): R$ {brl(res['preco_estimado'], 4)}**")
    else:
        L.append("Não foi possível estimar o preço (sem preços válidos).")
    L.append("")
    if res["avisos"]:
        L += ["## 5. Pontos de atenção", ""] + [f"- {x}" for x in res["avisos"]] + [""]
    if res["excluidos"]:
        L += ["## 6. Registros desconsiderados e motivo", "", "| Linha | Preço | Fornecedor | Órgão | Motivo |", "|---|---|---|---|---|"]
        for r in res["excluidos"]:
            L.append(f"| {r.get('_linha')} | {brl(numero(r.get(p['campo_preco'])), 4)} | {r.get('fornecedor') or ''} | {r.get('orgao') or ''} | {r.get('_excluido')} |")
        L.append("")
    return "\n".join(L)


def numero(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def justificativa_md(res, titulo):
    p, f = res["parametros"], res["estatisticas_finais"]
    cv = f.get("coef_variacao")
    corte = ("Não foi aplicado corte estatístico." if p["outliers"] == "nenhum"
             else f"Foram desconsiderados valores inconsistentes, inexequíveis ou excessivamente elevados segundo o critério fundamentado a seguir: {res['criterio_corte']}.")
    pouco = ""
    if f.get("n", 0) < p["min_precos"]:
        pouco = ("\n> **ATENÇÃO:** menos de três preços válidos. A formação do preço com menos de três preços é excepcional e exige "
                 "justificativa nos autos e aprovação da autoridade competente (IN SEGES/ME nº 65/2021, art. 6º, §5º — CONFERIR vigência/norma local). "
                 "Complete aqui a justificativa: [MOTIVO].\n")
    return f"""# Justificativa da Pesquisa de Preços — {titulo}

**Órgão/entidade:** [NOME DO ÓRGÃO] · **Processo:** [Nº DO PROCESSO] · **Data da pesquisa:** {date.today().strftime('%d/%m/%Y')}
**Responsável(is) pela pesquisa:** [NOME, CARGO/MATRÍCULA]

## 1. Objeto
[DESCRIÇÃO DO OBJETO, com código CATMAT/CATSER: ___]

## 2. Base normativa
Art. 23 da Lei nº 14.133/2021 e [IN SEGES/ME nº 65/2021, se órgão federal; ou o regulamento do seu ente — CONFERIR].

## 3. Parâmetros e fontes utilizados
Foi utilizado o parâmetro de **contratações similares feitas pela Administração Pública** (art. 23, §1º, II, da Lei nº 14.133/2021), com preços unitários **homologados** extraídos do Portal Nacional de Contratações Públicas (PNCP) e do Portal de Dados Abertos do Compras.gov.br, em contratações cujo resultado data de, no máximo, [JANELA: 365] dias anteriores à pesquisa (confira o --dias usado em buscar_precos.py). Cada preço consta do Mapa de Preços (anexo) com o número de controle e o link da contratação no PNCP.
[Se outros parâmetros foram usados — mídia especializada, pesquisa direta com fornecedores, notas fiscais — descreva-os aqui. Se foi usado somente este parâmetro, justifique: ___]

## 4. Série de preços e saneamento
Foram coletados {res['entrada_total']} registros; após filtros objetivos ([DESCREVA: quantidade comparável, UF/região, especificação]) e remoção de duplicatas, restaram {res['apos_filtros_e_duplicatas']}. {corte} Restaram **{f.get('n', 0)} preços válidos**.{pouco}

## 5. Método estatístico
Preço estimado pela **{p['metodo']}** dos preços válidos, na forma do art. 6º da IN SEGES/ME nº 65/2021 (CONFERIR). Resultado: **R$ {brl(res['preco_estimado'], 2) if res['preco_estimado'] is not None else '-'}** por unidade{f" (coeficiente de variação {cv:.1%})" if cv is not None else ""}.
[Justifique a escolha do método. Ex.: mediana por ser menos sensível a valores extremos.]
[Se aplicou acréscimo/decréscimo percentual (art. 6º §2º) ou fator de atualização ({p['fator_atualizacao']:g}), explique: ___]

## 6. Análise crítica
[Comente a dispersão, as diferenças de condições comerciais (prazo, local de entrega, quantidade, garantia) e os alertas de qualidade de dados do mapa. Pontos de atenção do cálculo:]
{chr(10).join('- ' + x for x in res['avisos']) or '- nenhum'}

## 7. Anexos
Mapa de Preços (`.mapa.csv/.xlsx`) · Memória de cálculo (`.memoria.md`) · [prints/relatórios do PNCP, se exigidos pelo órgão]

[LOCAL], {date.today().strftime('%d/%m/%Y')} — [ASSINATURA DO RESPONSÁVEL]
"""


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("calculo", help="arquivo .calculo.json gerado por calcular_precos.py")
    ap.add_argument("--titulo", default="[OBJETO DA CONTRATAÇÃO]", help="título/objeto para os cabeçalhos")
    ap.add_argument("--out", required=True, help="prefixo dos arquivos de saída")
    a = ap.parse_args()
    with open(a.calculo, encoding="utf-8") as f:
        res = json.load(f)
    linhas = linhas_mapa(res)
    grava_csv(linhas, a.out + ".mapa.csv")
    xlsx = grava_xlsx(linhas, res, a.out + ".mapa.xlsx")
    with open(a.out + ".memoria.md", "w", encoding="utf-8") as f:
        f.write(memoria_md(res, a.titulo))
    with open(a.out + ".justificativa.md", "w", encoding="utf-8") as f:
        f.write(justificativa_md(res, a.titulo))
    print(f"Gravados: {a.out}.mapa.csv, {'.mapa.xlsx, ' if xlsx else ''}.memoria.md, .justificativa.md", file=sys.stderr)


if __name__ == "__main__":
    main()
