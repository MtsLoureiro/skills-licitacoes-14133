#!/usr/bin/env python3
"""Monta o orçamento de obra/serviço de engenharia: quantitativos -> custo direto (SINAPI) -> BDI -> preço.

Entrada: planilha de quantitativos (CSV ; ou ,  ou XLSX simples) com as colunas
  item        numeração (1, 1.1, ...)                      [opcional]
  fonte       SINAPI | SICRO | PROPRIA | COTACAO | MANUAL   [padrão SINAPI]
  codigo      código SINAPI (composição ou insumo)         [obrigatório se fonte=SINAPI]
  descricao   texto (obrigatório se não for SINAPI)
  unidade     (obrigatória se não for SINAPI)
  quantidade  número
  custo_unitario   custo direto unitário SEM BDI (obrigatório se não for SINAPI; se informado em SINAPI, SOBRESCREVE e gera alerta)
  bdi         BDI % desta linha (opcional; senão vale o BDI global). Use BDI reduzido para fornecimento de materiais/equipamentos
              específicos quando cabível (art. 9º §1º do Decreto 7.983/2013; Súmula TCU 253)
  etapa       grupo/etapa para subtotais                    [opcional]
  referencia  origem do valor (composição analítica própria, nº das cotações, tabela e data-base)  [recomendado]
  tipo_sinapi comp | insumo (só se o código existir nos dois)  [opcional]

Saídas (--out PREFIXO): .orcamento.csv, .orcamento.xlsx (se openpyxl instalado), .abc.csv, .memoria.md, .justificativa.md, .orcamento.json

Exemplo:
  orcamento.py quantitativos.csv --uf DF --regime nao-desonerado --bdi 24.0 --out obra/orc

Regra de ouro: regime (desonerado/não) do SINAPI e do BDI (CPRB) têm de combinar. A data-base (mês do SINAPI) vai no documento.
"""
import argparse
import csv
import json
import math
import os
import sys
from datetime import date

import _sinapi as S
from curva_abc import curva_abc


def ler_quantitativos(caminho):
    if caminho.lower().endswith(".xlsx"):
        x = S.Xlsx(caminho)
        aba = next(iter(x.abas))
        linhas = [r for r in x.linhas(aba)]
        cab = None
        out = []
        for r in linhas:
            vals = {c: v for c, v in r.items() if c not in ("_n", "_f")}
            if cab is None:
                cab = {c: str(v).strip().lower() for c, v in vals.items()}
                continue
            out.append({cab[c]: v for c, v in vals.items() if c in cab})
        return out
    with open(caminho, encoding="utf-8-sig", newline="") as f:
        amostra = f.read(2048)
        f.seek(0)
        delim = ";" if amostra.count(";") >= amostra.count(",") else ","
        return [{(k or "").strip().lower(): (v or "").strip() for k, v in r.items()} for r in csv.DictReader(f, delimiter=delim)]


def numero(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("R$", "").strip()
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def arred(v, modo):
    """Duas casas. 'arredondar' (padrão) ou 'truncar'. Escolha e registre no edital (CONFERIR a regra do seu órgão)."""
    if modo == "truncar":
        return math.floor(v * 100 + 1e-9) / 100
    return round(v + 1e-12, 2)


def buscar_sinapi(con, regime, uf, codigo, tipo=None):
    achados = []
    if tipo in (None, "comp"):
        r = con.execute("""select d.descricao, d.unidade, c.custo, c.pct_as from composicao_def d join composicao c
            on c.regime=d.regime and c.codigo=d.codigo where d.regime=? and c.uf=? and d.codigo=?""", (regime, uf, codigo)).fetchone()
        if r:
            achados.append(("composição", r["descricao"], r["unidade"], r["custo"], r["pct_as"] or 0.0))
    if tipo in (None, "insumo"):
        r = con.execute("""select d.descricao, d.unidade, i.preco from insumo_def d join insumo i
            on i.regime=d.regime and i.codigo=d.codigo where d.regime=? and i.uf=? and d.codigo=?""", (regime, uf, codigo)).fetchone()
        if r:
            achados.append(("insumo", r["descricao"], r["unidade"], r["preco"], 0.0))
    return achados


def montar(linhas, con, args):
    itens, erros = [], []
    for n, l in enumerate(linhas, 1):
        fonte = (l.get("fonte") or "SINAPI").strip().upper()
        qtd = numero(l.get("quantidade"))
        alertas = []
        if qtd is None or qtd < 0:
            erros.append(f"linha {n}: quantidade inválida ({l.get('quantidade')!r})")
            continue
        it = {"n": n, "item": l.get("item") or str(n), "etapa": l.get("etapa") or "", "fonte": fonte, "codigo": l.get("codigo") or "",
              "quantidade": qtd, "referencia": l.get("referencia") or "", "pct_as": 0.0}
        manual = numero(l.get("custo_unitario"))
        if fonte == "SINAPI":
            try:
                cod = int(float(l.get("codigo")))
            except (TypeError, ValueError):
                erros.append(f"linha {n}: fonte SINAPI sem código válido")
                continue
            ach = buscar_sinapi(con, args.regime, args.uf, cod, (l.get("tipo_sinapi") or "").lower() or None)
            if not ach:
                if manual is None:
                    erros.append(f"linha {n}: código SINAPI {cod} não existe em {args.uf}/{args.regime} nesta data-base (não invento valor)")
                    continue
                alertas.append(f"código SINAPI {cod} não encontrado; usado custo informado")
                it.update(tipo_sinapi="-", descricao=l.get("descricao", ""), unidade=l.get("unidade", ""), custo=manual)
            else:
                if len(ach) > 1:
                    alertas.append(f"código {cod} existe como composição e insumo; usada a composição (use tipo_sinapi para forçar)")
                tipo, desc, un, custo, pas = ach[0]
                if not custo or custo <= 0:
                    if manual is None:
                        erros.append(f"linha {n}: código SINAPI {cod} ({desc[:50]}) SEM PREÇO em {args.uf}/{args.regime}: informe custo_unitario com a fonte")
                        continue
                    alertas.append("SINAPI sem preço na UF; usado custo informado")
                    custo = manual
                elif manual is not None and abs(manual - custo) > 0.005:
                    alertas.append(f"custo informado ({manual:.2f}) sobrescreve o SINAPI ({custo:.2f}): justificar")
                    custo = manual
                it.update(tipo_sinapi=tipo, descricao=desc, unidade=un, custo=custo, pct_as=pas)
                if pas and pas > 0:
                    alertas.append(f"%AS={pas:.1%}: parte do preço atribuída a partir de SP")
                if l.get("unidade") and S.normaliza(l["unidade"]) != S.normaliza(un):
                    alertas.append(f"unidade informada ({l['unidade']}) difere da SINAPI ({un}): conferir")
        else:
            if manual is None or manual <= 0:
                erros.append(f"linha {n}: fonte {fonte} exige custo_unitario > 0")
                continue
            if not l.get("descricao") or not l.get("unidade"):
                erros.append(f"linha {n}: fonte {fonte} exige descricao e unidade")
                continue
            it.update(tipo_sinapi="-", descricao=l["descricao"], unidade=l["unidade"], custo=manual)
            dica = {"PROPRIA": "composição própria: anexar a composição analítica (coeficientes e preços de insumos com fonte e data-base) e justificar por que o SINAPI/SICRO não atende (Decreto 7.983/2013, art. 6º — CONFERIR)",
                    "COTACAO": "cotação de mercado: anexar propostas formais (mínimo 3 sugerido; CONFERIR a norma do ente) e justificar a vantagem sobre o sistema oficial (Decreto 7.983/2013, art. 8º — CONFERIR)",
                    "SICRO": "SICRO informado manualmente: registrar tabela, UF/região, mês e regime",
                    "MANUAL": "valor manual: registrar a origem"}.get(fonte, "fonte não padrão: registrar a origem")
            if not it["referencia"]:
                alertas.append(dica)
        bdi = numero(l.get("bdi"))
        it["bdi"] = args.bdi if bdi is None else bdi
        if bdi is not None and bdi != args.bdi:
            alertas.append(f"BDI diferenciado ({bdi:g}%): justificar (ex.: fornecimento de materiais/equipamentos específicos)")
        it["custo"] = arred(it["custo"], args.arredondamento)
        it["preco_unitario"] = arred(it["custo"] * (1 + it["bdi"] / 100.0), args.arredondamento)
        it["custo_total"] = arred(it["custo"] * qtd, args.arredondamento)
        it["total"] = arred(it["preco_unitario"] * qtd, args.arredondamento)
        it["alertas"] = alertas
        itens.append(it)
    return itens, erros


def grava_csv(itens, caminho):
    cols = ["item", "etapa", "fonte", "codigo", "tipo_sinapi", "descricao", "unidade", "quantidade", "custo", "bdi", "preco_unitario", "custo_total", "total", "referencia", "alertas"]
    with open(caminho, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(cols)
        for i in itens:
            w.writerow([(" | ".join(i[c]) if c == "alertas" else i.get(c, "")) for c in cols])


def grava_xlsx(itens, resumo, caminho):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("[info] openpyxl não instalado: pulei o .xlsx (pip install openpyxl); o .csv tem os mesmos dados", file=sys.stderr)
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Orçamento"
    cab = ["Item", "Etapa", "Fonte", "Código", "Descrição", "Un.", "Quant.", "Custo unit. (R$)", "BDI (%)", "Preço unit. (R$)", "Total (R$)", "Alertas"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for k, i in enumerate(itens, 2):
        ws.append([i["item"], i["etapa"], i["fonte"], i["codigo"], i["descricao"], i["unidade"], i["quantidade"], i["custo"], i["bdi"],
                   f"=ROUND(H{k}*(1+I{k}/100),2)" if resumo["arredondamento"] == "arredondar" else i["preco_unitario"],
                   f"=ROUND(J{k}*G{k},2)" if resumo["arredondamento"] == "arredondar" else i["total"], " | ".join(i["alertas"])])
    ult = len(itens) + 1
    ws.append([])
    ws.append(["", "", "", "", "TOTAL (preço global de referência)", "", "", "", "", "", f"=SUM(K2:K{ult})"])
    for col, larg in zip("ABCDEFGHIJKL", (7, 14, 9, 9, 70, 6, 10, 14, 8, 14, 16, 60)):
        ws.column_dimensions[col].width = larg
    w2 = wb.create_sheet("Resumo")
    for k, v in resumo.items():
        w2.append([k, v if not isinstance(v, (dict, list)) else json.dumps(v, ensure_ascii=False)])
    wb.save(caminho)
    return True


def memoria_md(itens, resumo, abc):
    L = [f"# Memória de cálculo do orçamento — {resumo['titulo']}", "",
         f"Gerada em {date.today():%d/%m/%Y} por `orcamento.py`.", "",
         "## Parâmetros", "",
         f"- Base de custos: **SINAPI {resumo['data_base']}** (referência {resumo['mes_referencia']}, emissão {resumo['emissao']}), UF **{resumo['uf']}**, regime **{resumo['regime']}**",
         f"- Encargos sociais sobre a mão de obra (referência do SINAPI, {resumo['uf']}/{resumo['regime']}): horista {resumo['encargos'].get('horista')}, mensalista {resumo['encargos'].get('mensalista')} (percentuais aplicados dentro das composições)",
         f"- BDI global: **{resumo['bdi']:g}%** ({resumo['bdi_origem']})",
         f"- Arredondamento: {resumo['arredondamento']} a 2 casas (custo unitário, preço unitário e totais por linha)",
         f"- Fórmula: preço unitário = custo unitário × (1 + BDI); total da linha = preço unitário × quantidade", "",
         "## Resultado", "",
         f"- Itens: {len(itens)} · Custo direto total: **R$ {resumo['custo_direto']:,.2f}** · Preço global de referência (com BDI): **R$ {resumo['preco_global']:,.2f}**",
         f"- Itens SINAPI: {resumo['n_sinapi']} · outras fontes: {resumo['n_outras']} (peso {resumo['peso_outras']:.1f}% do preço)", ""]
    a = sum(1 for x in abc if x["classe"] == "A")
    L += [f"- Curva ABC: {a} itens classe A (até {resumo['corte_a']:g}% do valor; convenção, não norma)", "",
          "## Pontos de atenção", ""]
    com = [i for i in itens if i["alertas"]]
    L += [f"- Item {i['item']} ({i['descricao'][:50]}): " + "; ".join(i["alertas"]) for i in com] or ["- nenhum"]
    L += ["", "## Verificações antes de assinar", "",
          "- Regime (desonerado ou não) do SINAPI combina com o BDI (CPRB)? ",
          "- Data-base do SINAPI compatível com a data do orçamento/edital? (tabela mensal; atualizar se passar o prazo definido no edital)",
          "- Itens com `%AS` > 0 ou sem preço na UF têm justificativa?",
          "- Composições próprias e cotações têm documento de suporte anexado?",
          "- Administração local, canteiro e mobilização/desmobilização estão discriminados na planilha (não escondidos no BDI)? (TCU, Acórdão 2622/2013-Plenário, determinação reproduzida no Acórdão 1799/2014 — CONFERIR)"]
    return "\n".join(L)


def justificativa_md(resumo):
    return f"""# Justificativa do orçamento de referência — {resumo['titulo']}

**Órgão:** [NOME] · **Processo:** [Nº] · **Responsável técnico (engenheiro/arquiteto, com ART/RRT):** [NOME, REGISTRO] · **Data:** {date.today():%d/%m/%Y}

## 1. Fundamento
Art. 23, § 2º, da Lei nº 14.133/2021 (obras e serviços de engenharia: composição de custos unitários menores ou iguais à **mediana** do item no SINAPI ou SICRO, acrescidos de BDI e encargos sociais).
[Órgão federal: IN SEGES/ME nº 91/2022, que autoriza aplicar o Decreto nº 7.983/2013. Outro ente: norma local — CONFERIR. Estado/Município sem recursos da União: art. 23 § 3º permite outros sistemas de custos do ente.]

## 2. Base de custos
SINAPI, data-base **{resumo['mes_referencia']}** (relatório emitido em {resumo['emissao']}), localidade **{resumo['uf']}**, regime de encargos **{resumo['regime']}**. Os custos unitários do SINAPI são preços medianos de insumos pesquisados pelo IBGE e composições mantidas pela Caixa (Decreto 7.983/2013, art. 3º, § 1º).
Fontes complementares usadas em {resumo['n_outras']} item(ns) ({resumo['peso_outras']:.1f}% do preço): [descrever composições próprias/cotações e a justificativa da impossibilidade de usar o SINAPI/SICRO — Decreto 7.983/2013, arts. 6º e 8º — CONFERIR].

## 3. BDI
Taxa global de **{resumo['bdi']:g}%** ({resumo['bdi_origem']}), com composição discriminada em anexo: administração central, seguro, risco, garantia, despesas financeiras, lucro e tributos incidentes sobre o preço (Decreto 7.983/2013, art. 9º).
[Confirmar: ISS pela alíquota do município; PIS/COFINS pelo regime; CPRB só se o orçamento é desonerado; faixa referencial do TCU para o tipo de obra — Acórdão 2622/2013-Plenário, CONFERIR.]
[Se houver BDI reduzido para fornecimento de materiais/equipamentos específicos: justificar — art. 9º § 1º do Decreto 7.983/2013 e Súmula TCU 253.]

## 4. Resultado
Custo direto: R$ {resumo['custo_direto']:,.2f}. Preço global de referência (com BDI): **R$ {resumo['preco_global']:,.2f}**. Planilha, memória de cálculo e curva ABC em anexo.

## 5. Critérios de aceitabilidade de preços
[Definir no edital preços máximos unitários e global (Decreto 7.983/2013, arts. 2º IX, 11 e 13 — CONFERIR a redação vigente); em empreitada por preço global/integral, também por etapa do cronograma físico-financeiro.]

## 6. Observações
[Regime de execução, projeto básico/executivo e ART do orçamento (Decreto 7.983/2013, art. 10); sigilo do orçamento (Lei 14.133, art. 24 — CONFERIR).]

[LOCAL], {date.today():%d/%m/%Y} — [ASSINATURA DO RESPONSÁVEL TÉCNICO]
"""


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("quantitativos")
    ap.add_argument("--uf", default="DF")
    ap.add_argument("--regime", choices=list(S.REGIMES), default="nao-desonerado")
    ap.add_argument("--bdi", type=float, required=True, help="BDI global em %% (calcule com bdi.py)")
    ap.add_argument("--bdi-origem", default="informado pelo usuário [descrever a composição]")
    ap.add_argument("--mes", help="AAAA-MM do SINAPI (padrão: mais recente disponível)")
    ap.add_argument("--arquivo", help="ZIP/XLSX local do SINAPI (fallback manual)")
    ap.add_argument("--arredondamento", choices=["arredondar", "truncar"], default="arredondar")
    ap.add_argument("--titulo", default="[OBJETO DA OBRA/SERVIÇO]")
    ap.add_argument("--corte-a", type=float, default=80.0)
    ap.add_argument("--corte-b", type=float, default=95.0)
    ap.add_argument("--out", help="prefixo dos arquivos de saída")
    args = ap.parse_args()
    args.uf = args.uf.upper()

    db, tag = S.garantir_base(args.mes, args.arquivo)
    con = S.abrir(db)
    meta = dict(con.execute("select chave, valor from meta").fetchall())
    enc = con.execute("select horista, mensalista from encargos where regime=? and uf=?", (args.regime, args.uf)).fetchone()
    itens, erros = montar(ler_quantitativos(args.quantitativos), con, args)
    for e in erros:
        print("ERRO:", e, file=sys.stderr)
    if erros:
        sys.exit(f"{len(erros)} erro(s): corrija a planilha de quantitativos (nada foi gerado)")
    cd = sum(i["custo_total"] for i in itens)
    pg = sum(i["total"] for i in itens)
    outras = [i for i in itens if i["fonte"] != "SINAPI"]
    resumo = {"titulo": args.titulo, "data_base": tag, "mes_referencia": meta.get("mes_referencia"), "emissao": meta.get("data_emissao"),
              "uf": args.uf, "regime": args.regime, "bdi": args.bdi, "bdi_origem": args.bdi_origem, "arredondamento": args.arredondamento,
              "encargos": dict(enc) if enc else {}, "custo_direto": cd, "preco_global": pg, "n_sinapi": len(itens) - len(outras),
              "n_outras": len(outras), "peso_outras": (sum(i["total"] for i in outras) / pg * 100 if pg else 0.0),
              "corte_a": args.corte_a, "corte_b": args.corte_b}
    abc = curva_abc([{"item": i["item"], "descricao": i["descricao"], "valor": i["total"]} for i in itens], args.corte_a, args.corte_b)

    print(f"SINAPI {tag} (ref. {meta.get('mes_referencia')}) · {args.uf} · {args.regime} · BDI {args.bdi:g}%")
    print(f"{'item':<6}{'cód':>8} {'un':<4}{'qtd':>10} {'custo':>10} {'preço un.':>10} {'total':>13}  descrição")
    for i in itens:
        print(f"{i['item']:<6}{i['codigo']:>8} {i['unidade']:<4}{i['quantidade']:>10,.2f} {i['custo']:>10,.2f} {i['preco_unitario']:>10,.2f} {i['total']:>13,.2f}  {i['descricao'][:55]}{'  [!]' if i['alertas'] else ''}")
    print(f"\nCUSTO DIRETO R$ {cd:,.2f}  |  PREÇO GLOBAL (BDI {args.bdi:g}%) R$ {pg:,.2f}")
    for i in itens:
        for al in i["alertas"]:
            print(f"  [!] item {i['item']}: {al}")
    if args.out:
        os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
        grava_csv(itens, args.out + ".orcamento.csv")
        x = grava_xlsx(itens, resumo, args.out + ".orcamento.xlsx")
        with open(args.out + ".abc.csv", "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["classe", "item", "descricao", "valor", "pct", "acumulado_pct"])
            for r in abc:
                w.writerow([r["classe"], r["item"], r["descricao"], round(r["valor"], 2), round(r["pct"], 4), round(r["acumulado"], 4)])
        open(args.out + ".memoria.md", "w", encoding="utf-8").write(memoria_md(itens, resumo, abc))
        open(args.out + ".justificativa.md", "w", encoding="utf-8").write(justificativa_md(resumo))
        json.dump({"resumo": resumo, "itens": itens}, open(args.out + ".orcamento.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print(f"Gravados: {args.out}.orcamento.csv{', .xlsx' if x else ''}, .abc.csv, .memoria.md, .justificativa.md, .orcamento.json", file=sys.stderr)


if __name__ == "__main__":
    main()
